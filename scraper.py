"""書籍一覧スクレイピングツール

練習用サイト books.toscrape.com から本の一覧（タイトル・価格・評価・在庫・URL）を集め、
CSV と Excel に書き出す。

使い方:
    python scraper.py              # 1ページ目（20冊）だけ取得
    python scraper.py --pages 5    # 5ページ（100冊）取得
    python scraper.py --pages 0    # 全ページ取得
"""

from __future__ import annotations

import argparse
import csv
import logging
import time
from dataclasses import asdict, dataclass, fields
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

START_URL = "https://books.toscrape.com/catalogue/page-1.html"
USER_AGENT = "book-scraper/1.0 (portfolio sample)"

# サイト上の評価はクラス名（"star-rating Three" など）で表されているので数値に直す
RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

# 出力ファイルの見出し（Book の項目と同じ順番）
HEADERS = {
    "title": "タイトル",
    "price_gbp": "価格（£）",
    "rating": "評価（5段階）",
    "in_stock": "在庫",
    "url": "URL",
}

log = logging.getLogger("book-scraper")


@dataclass
class Book:
    title: str
    price_gbp: float
    rating: int
    in_stock: str
    url: str


def create_session() -> requests.Session:
    """一時的なエラー（503 など）のときに、間隔をあけて自動で再試行するセッションを作る。"""
    retry = Retry(total=3, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504])
    session = requests.Session()
    session.mount("https://", HTTPAdapter(max_retries=retry))
    session.headers["User-Agent"] = USER_AGENT
    return session


def fetch_html(session: requests.Session, url: str) -> BeautifulSoup:
    response = session.get(url, timeout=10)
    response.raise_for_status()
    # サイトは UTF-8 だが requests が文字コードを誤判定し、「£」が文字化けするので明示する
    response.encoding = "utf-8"
    return BeautifulSoup(response.text, "html.parser")


def parse_books(soup: BeautifulSoup, page_url: str) -> list[Book]:
    """一覧ページ1枚分から本の情報を取り出す。"""
    books = []
    for item in soup.select("article.product_pod"):
        link = item.select_one("h3 a")
        price_text = item.select_one("p.price_color").get_text(strip=True)
        rating_classes = item.select_one("p.star-rating")["class"]

        books.append(
            Book(
                # 表示上のタイトルは「...」で省略されるため、title 属性から正式名を取る
                title=link["title"],
                price_gbp=float(price_text.lstrip("£")),
                rating=next((RATING_WORDS[c] for c in rating_classes if c in RATING_WORDS), 0),
                in_stock="あり" if "In stock" in item.select_one("p.availability").get_text() else "なし",
                url=urljoin(page_url, link["href"]),
            )
        )
    return books


def find_next_page(soup: BeautifulSoup, page_url: str) -> str | None:
    next_link = soup.select_one("li.next a")
    return urljoin(page_url, next_link["href"]) if next_link else None


def scrape(max_pages: int, delay: float) -> list[Book]:
    """一覧ページを順にたどって本を集める。max_pages が 0 なら最後のページまで。"""
    session = create_session()
    books: list[Book] = []
    url: str | None = START_URL
    page = 0

    while url and (max_pages == 0 or page < max_pages):
        page += 1
        log.info("%dページ目を取得中: %s", page, url)
        soup = fetch_html(session, url)
        books.extend(parse_books(soup, url))
        url = find_next_page(soup, url)

        # サーバーに負担をかけないよう、次のページの前に少し待つ
        if url and (max_pages == 0 or page < max_pages):
            time.sleep(delay)

    return books


def save_csv(books: list[Book], path: Path) -> None:
    # utf-8-sig にすると、Excel で開いても日本語の見出しが文字化けしない
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(HEADERS))
        writer.writerow(HEADERS)
        writer.writerows(asdict(book) for book in books)


def save_excel(books: list[Book], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "書籍一覧"

    ws.append(list(HEADERS.values()))
    for book in books:
        ws.append([getattr(book, f.name) for f in fields(Book)])

    # 見出しを目立たせ、スクロールしても見出しが残るようにし、絞り込みを付ける
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F4F4F")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "£#,##0.00"

    widths = {"title": 60, "price_gbp": 12, "rating": 14, "in_stock": 8, "url": 70}
    for i, key in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[key]

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="books.toscrape.com の書籍一覧を CSV / Excel に書き出す")
    parser.add_argument("--pages", type=int, default=1, help="取得するページ数（0 で全ページ。1ページ20冊）")
    parser.add_argument("--delay", type=float, default=1.0, help="ページ取得の間隔（秒）")
    parser.add_argument("--out-dir", type=Path, default=Path("output"), help="出力先フォルダ")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s", datefmt="%H:%M:%S")

    books = scrape(args.pages, args.delay)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_path = args.out_dir / f"books_{stamp}.csv"
    xlsx_path = args.out_dir / f"books_{stamp}.xlsx"
    save_csv(books, csv_path)
    save_excel(books, xlsx_path)

    log.info("%d冊を保存しました: %s / %s", len(books), csv_path, xlsx_path)


if __name__ == "__main__":
    main()
